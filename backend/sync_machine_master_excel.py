import argparse
import csv
from pathlib import Path
from openpyxl import load_workbook


def norm(v):
    return "" if v is None else str(v).strip()


def main():
    parser = argparse.ArgumentParser(
        description="Append approved NEW_MACHINE codes to MC NAME and finalize machine resolution mapping."
    )
    parser.add_argument("xlsx", help="Path to ROD MM STOCK.xlsx")
    parser.add_argument(
        "resolution_csv",
        help="Path to machine_resolution_review.csv",
    )
    parser.add_argument(
        "--output-xlsx",
        default="",
        help="Output workbook path. Default: <input>_MC_UPDATED.xlsx",
    )
    parser.add_argument(
        "--output-mapping",
        default="",
        help="Output finalized mapping CSV. Default: machine_resolution_final.csv beside resolution CSV",
    )
    args = parser.parse_args()

    xlsx_path = Path(args.xlsx)
    resolution_path = Path(args.resolution_csv)

    if not xlsx_path.exists():
        raise SystemExit(f"Excel not found: {xlsx_path}")
    if not resolution_path.exists():
        raise SystemExit(f"Resolution CSV not found: {resolution_path}")

    output_xlsx = (
        Path(args.output_xlsx)
        if args.output_xlsx
        else xlsx_path.with_name(xlsx_path.stem + "_MC_UPDATED.xlsx")
    )
    output_mapping = (
        Path(args.output_mapping)
        if args.output_mapping
        else resolution_path.with_name("machine_resolution_final.csv")
    )

    # Read mapping/review file.
    with resolution_path.open("r", newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        rows = list(reader)
        fields = list(reader.fieldnames or [])

    if "reference_code" not in fields or "decision" not in fields:
        raise SystemExit("Resolution CSV must contain reference_code and decision columns.")

    if "result_code_type" not in fields:
        fields.append("result_code_type")

    new_machine_codes = []
    decision_counts = {}

    for row in rows:
        code = norm(row.get("reference_code"))
        decision = norm(row.get("decision")).upper()

        if not code:
            continue

        decision_counts[decision or "BLANK"] = decision_counts.get(decision or "BLANK", 0) + 1

        if decision == "NEW_MACHINE":
            new_machine_codes.append(code)
            row["result_code_type"] = "CURRENT"
        elif decision == "OLD":
            row["result_code_type"] = "OLD"
        elif decision == "ALIAS":
            row["result_code_type"] = "ALIAS"
        elif decision in {"LINE", "KEEP_REFERENCE", "IGNORE", ""}:
            row["result_code_type"] = "REFERENCE"

    # Update MC NAME.
    wb = load_workbook(xlsx_path)
    if "MC NAME" not in wb.sheetnames:
        raise SystemExit(f"'MC NAME' sheet not found. Sheets: {wb.sheetnames}")

    ws = wb["MC NAME"]

    header = norm(ws.cell(row=1, column=1).value)
    if header.upper() != "MACHINE NAME":
        raise SystemExit(
            f"Expected MC NAME!A1 to be 'MACHINE NAME', got: {header!r}"
        )

    existing = {}
    for r in range(2, ws.max_row + 1):
        value = norm(ws.cell(row=r, column=1).value)
        if value:
            existing[value.upper()] = value

    added = []
    skipped = []

    for code in sorted(set(new_machine_codes), key=lambda x: x.upper()):
        if code.upper() in existing:
            skipped.append(code)
            continue

        ws.append([code])
        existing[code.upper()] = code
        added.append(code)

    # Save to a new workbook so the original stays untouched.
    wb.save(output_xlsx)

    # Save finalized mapping registry.
    with output_mapping.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)

    print("=" * 60)
    print("MC NAME + MAPPING SYNC COMPLETE")
    print("=" * 60)
    print("Input Excel              :", xlsx_path)
    print("Output Excel             :", output_xlsx)
    print("Input mapping            :", resolution_path)
    print("Output mapping           :", output_mapping)
    print()
    print("Approved NEW_MACHINE     :", len(set(new_machine_codes)))
    print("Added to MC NAME         :", len(added))
    print("Already in MC NAME       :", len(skipped))
    print("MC NAME unique after sync:", len(existing))
    print()
    for key in sorted(decision_counts):
        print(f"{key:20}: {decision_counts[key]}")
    print()
    print("Original workbook was NOT overwritten.")


if __name__ == "__main__":
    main()
