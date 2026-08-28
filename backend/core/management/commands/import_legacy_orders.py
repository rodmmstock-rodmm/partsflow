from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.db.models import Q
from django.utils import timezone
from openpyxl import load_workbook

from core.models import Employee, Machine, OrderRecord, Part, Supplier


def text(value):
    return "" if value is None else str(value).strip()


def as_decimal(value):
    try:
        return Decimal(str(value if value not in (None, "") else 0).replace(",", ""))
    except (InvalidOperation, TypeError, ValueError):
        return Decimal("0")


def as_date(value):
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raw = text(value)
    for fmt in ["%d/%m/%Y", "%Y-%m-%d", "%m/%d/%Y"]:
        try:
            return datetime.strptime(raw, fmt).date()
        except ValueError:
            pass
    return None


def truthy(value):
    return text(value).lower() in {"true", "1", "yes", "y", "cancel", "cancelled", "ยกเลิก"}


def rows(ws):
    values = list(ws.values)
    if not values:
        return []
    headers = [text(x) for x in values[0]]
    result = []
    for excel_row, row in enumerate(values[1:], start=2):
        if not any(x is not None for x in row):
            continue
        item = dict(zip(headers, row))
        item["_excel_row"] = excel_row
        result.append(item)
    return result


class Command(BaseCommand):
    help = "Import legacy ORDER sheet into OrderRecord. Does NOT modify Inventory."

    def add_arguments(self, parser):
        parser.add_argument("xlsx")
        parser.add_argument("--dry-run", action="store_true")

    @transaction.atomic
    def handle(self, *args, **options):
        path = options["xlsx"]
        dry_run = options["dry_run"]
        try:
            wb = load_workbook(path, data_only=True)
        except Exception as exc:
            raise CommandError(f"เปิด Excel ไม่สำเร็จ: {exc}")
        if "ORDER" not in wb.sheetnames:
            raise CommandError("ไม่พบ Sheet ORDER")

        imported = 0
        duplicates = 0
        skipped = 0
        warnings = []

        for row in rows(wb["ORDER"]):
            excel_row = row["_excel_row"]
            order_number = text(row.get("ORDER NUMBER")) or f"LEGACY-ORDER-{excel_row}"
            if OrderRecord.objects.filter(order_number=order_number).exists():
                duplicates += 1
                continue

            order_date = as_date(row.get("DATE")) or timezone.localdate()
            machine_text = text(row.get("MACHINE NAME"))
            machine = Machine.objects.filter(
                Q(code__iexact=machine_text) | Q(name__iexact=machine_text)
            ).first()
            if machine_text and not machine:
                warnings.append(f"row {excel_row}: machine '{machine_text}' not found")

            sku = text(row.get("ITEM ID"))
            part = Part.objects.select_related("maker", "unit").filter(sku__iexact=sku).first() if sku else None
            part_name = part.name if part else text(row.get("PART NAME"))
            part_detail = part.description if part else text(row.get("PART DETAIL"))
            maker = part.maker.name if part and part.maker else text(row.get("MAKER"))
            unit = part.unit.code if part and part.unit else text(row.get("UNIT"))
            if not part_name or not part_detail or not maker or not unit:
                skipped += 1
                warnings.append(f"row {excel_row}: required part text incomplete")
                continue

            ordered_text = text(row.get("ORDERED BY"))
            ordered_by = Employee.objects.filter(
                Q(employee_code__iexact=ordered_text) | Q(name__iexact=ordered_text)
            ).first() if ordered_text else None
            pic_text = text(row.get("PERSON IN CHARGE OF ORDER"))
            pic = Employee.objects.filter(
                Q(employee_code__iexact=pic_text) | Q(name__iexact=pic_text)
            ).first() if pic_text else None
            vendor_text = text(row.get("VENDOR ORDER"))
            vendor = Supplier.objects.filter(
                Q(code__iexact=vendor_text) | Q(name__iexact=vendor_text)
            ).first() if vendor_text else None

            amount = max(1, int(as_decimal(row.get("AMOUNT")) or 1))
            price_per_unit = as_decimal(row.get("PRICE PER UNIT"))
            price_total = as_decimal(row.get("PRICE TOTAL")) or Decimal(amount) * price_per_unit
            quotation = text(row.get("QUOTATION"))
            po_number = text(row.get("PO NUMBER"))
            issue_pr = as_date(row.get("ISSUE PR DATE"))
            due_date = as_date(row.get("DUE DATE"))
            vendor_confirm = as_date(row.get("VENDOR CONFIRM DATE"))
            receive_date = as_date(row.get("RECIEVE DATE"))
            received_at = None
            if receive_date:
                received_at = timezone.make_aware(datetime.combine(receive_date, time(12, 0)))

            po_group = [bool(po_number), bool(issue_pr), bool(due_date)]
            if any(po_group) and not all(po_group):
                warnings.append(f"row {excel_row}: PO/ISSUE PR/DUE incomplete; imported as Wait Quotation")
                po_number = ""
                issue_pr = None
                due_date = None

            if received_at:
                status = OrderRecord.STATUS_COMPLETE
            elif po_number and issue_pr and due_date:
                status = OrderRecord.STATUS_ITEM
            elif quotation:
                status = OrderRecord.STATUS_QUOTE
            else:
                status = OrderRecord.STATUS_NEW

            pending_date = as_date(row.get("วันที่ค้างดาต้า") or row.get("วันที่งานค้าง"))
            urgent_status = text(row.get("สถานะงานด่วน"))
            edit_status = text(row.get("EDIT DATA STATUS") or row.get("สถานะการแก้ไข data"))
            if pending_date and not edit_status:
                edit_status = OrderRecord.EDIT_WAIT_QUOTE

            group_order = text(row.get("GROUP ORDER"))
            if not group_order and machine and (urgent_status or pending_date):
                group_order = f"{machine.code}_{order_date.strftime('%d%m%Y')}"

            OrderRecord.objects.create(
                order_number=order_number,
                order_date=order_date,
                factory=text(row.get("FACTORY")) or "MM-4",
                group_order=group_order,
                machine=machine,
                job=text(row.get("JOB")).upper(),
                urgent_status=urgent_status,
                pending_data_date=pending_date,
                remark=text(row.get("REMARK")),
                quotation=quotation,
                part=part,
                part_name=part_name,
                part_detail=part_detail,
                maker_text=maker,
                amount=amount,
                unit_text=unit,
                po_number=po_number,
                price_per_unit=price_per_unit,
                price_total=price_total,
                vendor=vendor,
                lead_time_days=int(as_decimal(row.get("LEAD TIME")) or 0) or None,
                ordered_by=ordered_by,
                issue_pr_date=issue_pr,
                due_date=due_date,
                vendor_confirm_date=vendor_confirm,
                received_at=received_at,
                person_in_charge=pic,
                status=status,
                cancel_status=truthy(row.get("CANCLE STATUS") or row.get("สถานะยกเลิก")),
                edit_data_status=edit_status,
                edit_workflow_enabled=bool(pending_date),
                source_type="NORMAL",
                stock_received=bool(received_at and part and text(row.get("JOB")).upper() == "SPARE"),
                legacy_source="ORDER",
                legacy_id=f"ORDER:{excel_row}",
            )
            imported += 1

        self.stdout.write(f"Imported: {imported}")
        self.stdout.write(f"Duplicates: {duplicates}")
        self.stdout.write(f"Skipped: {skipped}")
        self.stdout.write(f"Warnings: {len(warnings)}")
        for warning in warnings[:50]:
            self.stdout.write(self.style.WARNING(warning))
        if len(warnings) > 50:
            self.stdout.write(self.style.WARNING(f"... and {len(warnings)-50} more warnings"))
        if dry_run:
            transaction.set_rollback(True)
            self.stdout.write(self.style.SUCCESS("DRY RUN: rolled back. Inventory was NOT modified."))
        else:
            self.stdout.write(self.style.SUCCESS("Order import complete. Inventory was NOT modified."))
