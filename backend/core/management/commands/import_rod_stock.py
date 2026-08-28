import csv
import hashlib
import re
from pathlib import Path
from collections import Counter, defaultdict
from decimal import Decimal, InvalidOperation

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from core.models import (
    Employee,
    Inventory,
    Location,
    Machine,
    MachineCode,
    Maker,
    Part,
    PartMachine,
    Supplier,
    Unit,
)


# ============================================================
# Helpers
# ============================================================

def text(value):
    return "" if value is None else str(value).strip()


def decimal_value(value):
    if value is None or value == "":
        return Decimal("0")
    try:
        return Decimal(str(value).replace(",", "").strip())
    except (InvalidOperation, ValueError, TypeError):
        return Decimal("0")


def positive_int(value):
    try:
        return max(0, int(decimal_value(value)))
    except Exception:
        return 0


def rows(workbook, sheet_name):
    if sheet_name not in workbook.sheetnames:
        raise CommandError(
            f"ไม่พบ Sheet '{sheet_name}' "
            f"(พบ: {', '.join(workbook.sheetnames)})"
        )

    ws = workbook[sheet_name]
    values = list(ws.values)

    if not values:
        return []

    headers = [text(x) for x in values[0]]
    result = []

    for excel_row, row in enumerate(values[1:], start=2):
        if not any(x is not None for x in row):
            continue

        data = dict(zip(headers, row))
        data["_excel_row"] = excel_row
        result.append(data)

    return result


def normalize_key(value):
    """
    Normalized only for comparison.
    Does NOT change what is stored in the database.
    """
    return re.sub(r"\s+", "", text(value)).upper()


def clean_machine_code(value):
    value = text(value)
    value = re.sub(r"\s+", "", value)
    return value.strip(" ,;")


def split_machine_codes(raw_value):
    """
    Parse Machine Name cells conservatively.

    Supports:
      BFD-13,15,17      -> BFD-13 / BFD-15 / BFD-17
      BSM-05.BSM-06     -> BSM-05 / BSM-06
      BSM-12CGM-07      -> BSM-12 / CGM-07
      CGM-10RLP-09      -> CGM-10 / RLP-09
      OVN-11BSM-10      -> OVN-11 / BSM-10

    It deliberately does NOT guess BFD-BUFF, BFG-NEW, BSM11B,
    TOOLING, etc. Those remain single references until classified.
    """
    raw = text(raw_value)
    if not raw:
        return []

    raw = raw.replace("\n", ",").replace("\r", ",").replace(";", ",")
    pieces = [clean_machine_code(x) for x in raw.split(",")]
    pieces = [x for x in pieces if x]

    result = []
    current_prefix = None

    def add_unique(value):
        value = clean_machine_code(value).strip(".:/")
        if value and value not in result:
            result.append(value)

    for item in pieces:
        # Detect two or more explicit PREFIX- starts inside one piece.
        # This fixes missing-delimiter data such as BSM-12CGM-07.
        prefix_matches = list(
            re.finditer(r"[A-Za-z][A-Za-z0-9]*-", item)
        )

        if len(prefix_matches) >= 2:
            extracted = []
            for idx, match in enumerate(prefix_matches):
                start = match.start()
                end = (
                    prefix_matches[idx + 1].start()
                    if idx + 1 < len(prefix_matches)
                    else len(item)
                )
                candidate = item[start:end].strip(".:/ ")

                if re.fullmatch(
                    r"[A-Za-z][A-Za-z0-9]*-\d+[A-Za-z]?",
                    candidate,
                ):
                    extracted.append(candidate)

            # Only split if all detected chunks are valid machine-like codes.
            # Otherwise preserve the original value as a reference.
            if len(extracted) == len(prefix_matches):
                for code_item in extracted:
                    add_unique(code_item)
                current_prefix = None
                continue

        # Numeric continuation, e.g. BFD-13,15,17
        if re.fullmatch(r"\d+", item) and current_prefix:
            item = current_prefix + item

        match = re.fullmatch(
            r"([A-Za-z][A-Za-z0-9]*-)(\d+[A-Za-z]?)",
            item,
        )
        if match:
            current_prefix = match.group(1).upper()
        elif re.search(r"[A-Za-z]", item):
            current_prefix = None

        add_unique(item)

    return result


def relocation_old_code(current_code):
    """
    Business rule confirmed by the user:
    machine codes ending in B are the renamed codes after relocation.

    Examples:
      BFD15B -> BFD-15
      IDT03B -> IDT-03
      MLG04B -> MLG-04

    Only derives an OLD alias from a CURRENT code in MC NAME.
    """
    code = clean_machine_code(current_code)
    match = re.fullmatch(r"([A-Za-z]+)-?(\d+)B", code, flags=re.IGNORECASE)
    if not match:
        return None
    return f"{match.group(1).upper()}-{match.group(2)}"


def supplier_code_base(name):
    slug = re.sub(r"[^A-Za-z0-9]+", "-", text(name)).strip("-").upper()
    if not slug:
        slug = "SUPPLIER"
    return ("V-" + slug)[:70]


def get_supplier(vendor_name):
    vendor_name = text(vendor_name)
    if not vendor_name:
        return None, False

    existing = Supplier.objects.filter(name__iexact=vendor_name).first()
    if existing:
        return existing, False

    base = supplier_code_base(vendor_name)
    candidate = base

    conflict = Supplier.objects.filter(code=candidate).first()
    if conflict and conflict.name.lower() != vendor_name.lower():
        suffix = hashlib.sha1(vendor_name.encode("utf-8")).hexdigest()[:8].upper()
        candidate = f"{base[:60]}-{suffix}"

    return Supplier.objects.get_or_create(
        code=candidate,
        defaults={"name": vendor_name},
    )


def load_confirmed_machine_mappings(csv_path):
    """
    Load business-approved machine-code mappings from CSV.

    Required columns:
      old_code,current_code

    Optional columns:
      code_type,note,active

    Example:
      BFD-15,BFD15B,OLD,Renamed after relocation,1

    This file is the source of truth for approved aliases.
    The importer never promotes a suggestion to an alias automatically.
    """
    mappings = []

    if not csv_path:
        return mappings

    path = Path(csv_path)
    if not path.exists():
        raise CommandError(f"Machine mapping CSV not found: {csv_path}")

    with path.open("r", newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)

        if not reader.fieldnames:
            raise CommandError(f"Machine mapping CSV has no header: {csv_path}")

        field_lookup = {
            normalize_key(name).replace("_", ""): name
            for name in reader.fieldnames
            if name
        }

        old_field = field_lookup.get("OLDCODE")
        current_field = field_lookup.get("CURRENTCODE")

        if not old_field or not current_field:
            raise CommandError(
                "Machine mapping CSV must contain columns: "
                "old_code,current_code"
            )

        seen_old_codes = set()

        for csv_row_number, row in enumerate(reader, start=2):
            old_code = clean_machine_code(row.get(old_field))
            current_code = clean_machine_code(row.get(current_field))

            # Allow blank lines.
            if not old_code and not current_code:
                continue

            if not old_code or not current_code:
                raise CommandError(
                    f"Machine mapping CSV row {csv_row_number}: "
                    "old_code/current_code is incomplete"
                )

            old_key = normalize_key(old_code)
            if old_key in seen_old_codes:
                raise CommandError(
                    f"Machine mapping CSV row {csv_row_number}: "
                    f"duplicate old_code: {old_code}"
                )
            seen_old_codes.add(old_key)

            code_type = text(row.get(field_lookup.get("CODETYPE", ""))).upper()
            if code_type not in {"OLD", "ALIAS"}:
                code_type = "OLD"

            active_raw = text(row.get(field_lookup.get("ACTIVE", ""))).lower()
            active = active_raw not in {"0", "false", "no", "n"}

            note = text(row.get(field_lookup.get("NOTE", "")))
            if not note:
                note = (
                    f"Business-approved mapping: {old_code} -> {current_code}"
                )

            mappings.append(
                {
                    "old_code": old_code,
                    "current_code": current_code,
                    "code_type": code_type,
                    "note": note,
                    "active": active,
                    "csv_row": csv_row_number,
                }
            )

    return mappings


def suggested_current_code(old_code, current_by_key):
    """
    Suggestion only. Never auto-applied.

    Example:
      BFD-15 -> BFD15B, if BFD15B exists in MC NAME.

    A human/business confirmation is still required before the mapping
    is added to the approved machine-mapping CSV.
    """
    code = clean_machine_code(old_code)
    match = re.fullmatch(r"([A-Za-z]+)-(\d+)", code, flags=re.IGNORECASE)
    if not match:
        return "", ""

    candidate = f"{match.group(1).upper()}{match.group(2)}B"
    machine = current_by_key.get(normalize_key(candidate))

    if machine:
        return machine.code, "possible old->new relocation code; CONFIRM before mapping"

    return "", ""



class Command(BaseCommand):
    help = "Import ROD MM STOCK master data with improved concatenated Machine Code parsing (V3.3)."

    def add_arguments(self, parser):
        parser.add_argument("xlsx", help="Path to ROD MM STOCK.xlsx")
        parser.add_argument(
            "--dry-run",
            action="store_true",
            help="Run all validation/import logic and rollback database changes.",
        )
        parser.add_argument(
            "--machine-mapping-csv",
            default="",
            help=(
                "CSV containing business-approved machine mappings "
                "(old_code,current_code,code_type,note,active)."
            ),
        )
        parser.add_argument(
            "--unmatched-limit",
            type=int,
            default=100,
            help="Maximum unmatched machine codes to print. Default: 100",
        )
        parser.add_argument(
            "--unmatched-csv",
            default="",
            help="Optional CSV path for the full unmatched machine report.",
        )

    @transaction.atomic
    def handle(self, *args, **options):
        xlsx = options["xlsx"]
        dry_run = options["dry_run"]
        unmatched_limit = max(0, options["unmatched_limit"])
        unmatched_csv = options["unmatched_csv"]
        machine_mapping_csv = options["machine_mapping_csv"]
        confirmed_mappings = load_confirmed_machine_mappings(
            machine_mapping_csv
        )

        self.stdout.write("")
        self.stdout.write(self.style.MIGRATE_HEADING("=" * 60))
        self.stdout.write(
            self.style.MIGRATE_HEADING(
                " PartsFlow - ROD MM STOCK Import V3.3"
            )
        )
        self.stdout.write(self.style.MIGRATE_HEADING("=" * 60))
        self.stdout.write(f"Excel: {xlsx}")
        if machine_mapping_csv:
            self.stdout.write(
                f"Machine mapping CSV: {machine_mapping_csv} "
                f"({len(confirmed_mappings)} approved mappings)"
            )
        else:
            self.stdout.write(
                self.style.WARNING(
                    "Machine mapping CSV: not supplied "
                    "(no old/alias mappings will be approved)"
                )
            )
        self.stdout.write("")

        try:
            workbook = load_workbook(xlsx, data_only=True)
        except Exception as exc:
            raise CommandError(f"ไม่สามารถเปิด Excel ได้: {exc}")

        self.stdout.write(f"Sheets: {', '.join(workbook.sheetnames)}")
        self.stdout.write("")

        employee_rows = rows(workbook, "NAME")
        machine_rows = rows(workbook, "MC NAME")
        stock_rows = rows(workbook, "STOCK")

        stats = Counter()
        warnings = []
        errors = []

        unresolved_occurrences = Counter()
        unresolved_parts = defaultdict(set)

        # ========================================================
        # EMPLOYEES
        # ========================================================
        self.stdout.write(f"Reading NAME: {len(employee_rows)} rows")

        for row in employee_rows:
            code = text(row.get("EMPLOYEE ID")) or text(row.get("NAME"))
            if not code:
                warnings.append(
                    f"NAME row {row['_excel_row']}: ไม่มี Employee ID / NAME"
                )
                continue

            try:
                with transaction.atomic():
                    _, created = Employee.objects.update_or_create(
                        employee_code=code,
                        defaults={
                            "name": text(row.get("NAME")) or code,
                            "role": text(row.get("ROLE")),
                            "legacy_source": "NAME",
                            "legacy_id": code,
                        },
                    )
                stats["employees_processed"] += 1
                stats["employees_created"] += int(created)
            except Exception as exc:
                errors.append(f"NAME row {row['_excel_row']}: {exc}")

        # ========================================================
        # MACHINE MASTER
        # ========================================================
        self.stdout.write(f"Reading MC NAME: {len(machine_rows)} rows")

        current_by_key = {}
        machine_row_codes = []
        current_code_duplicates = Counter()

        for row in machine_rows:
            code = (
                text(row.get("MACHINE NAME"))
                or text(row.get("MC"))
                or text(row.get("MACHINE"))
                or text(row.get("MC NAME"))
            )
            code = clean_machine_code(code)

            if not code:
                warnings.append(
                    f"MC NAME row {row['_excel_row']}: ไม่มี Machine Code"
                )
                continue

            machine_row_codes.append(code)
            current_code_duplicates[normalize_key(code)] += 1

            try:
                with transaction.atomic():
                    machine, created = Machine.objects.update_or_create(
                        code=code,
                        defaults={
                            "name": code,
                            "legacy_source": "MC NAME",
                            "legacy_id": code,
                        },
                    )

                    MachineCode.objects.update_or_create(
                        code=code,
                        defaults={
                            "machine": machine,
                            "code_type": "CURRENT",
                            "note": "Current code imported from MC NAME",
                            "active": True,
                        },
                    )

                current_by_key[normalize_key(code)] = machine
                stats["machine_rows_processed"] += 1
                stats["machines_created"] += int(created)
            except Exception as exc:
                errors.append(
                    f"MC NAME row {row['_excel_row']}: {code}: {exc}"
                )

        duplicate_machine_codes = [
            key for key, count in current_code_duplicates.items() if count > 1
        ]
        stats["current_machine_unique"] = len(current_by_key)
        stats["machine_duplicate_rows"] = sum(
            current_code_duplicates[k] - 1 for k in duplicate_machine_codes
        )

        # ========================================================
        # BUILD CONFIRMED ALIASES ONLY
        # ========================================================
        # IMPORTANT:
        # Do not automatically assume that every code ending in "B"
        # is the same physical machine as PREFIX-##.
        # Only aliases explicitly confirmed by the business are applied.
        alias_target_by_key = {}

        mapping_meta_by_key = {}

        for mapping in confirmed_mappings:
            old_code = mapping["old_code"]
            current_code = mapping["current_code"]
            target = current_by_key.get(normalize_key(current_code))

            if not target:
                warnings.append(
                    f"Confirmed mapping target not found in MC NAME "
                    f"(CSV row {mapping['csv_row']}): "
                    f"{old_code} -> {current_code}"
                )
                continue

            old_key = normalize_key(old_code)

            if (
                old_key in current_by_key
                and current_by_key[old_key].pk != target.pk
            ):
                warnings.append(
                    f"Confirmed mapping conflict "
                    f"(CSV row {mapping['csv_row']}): "
                    f"{old_code} -> {current_code}, "
                    f"but {old_code} is also a CURRENT machine."
                )
                continue

            alias_target_by_key[old_key] = target
            mapping_meta_by_key[old_key] = mapping

        stats["machine_mappings_loaded"] = len(confirmed_mappings)
        stats["confirmed_aliases_available"] = len(alias_target_by_key)

        # ========================================================
        # STOCK MASTER
        # ========================================================
        self.stdout.write(f"Reading STOCK: {len(stock_rows)} rows")
        self.stdout.write("")

        seen_skus = set()

        for row in stock_rows:
            excel_row = row["_excel_row"]
            sku = text(row.get("ITEM ID"))

            if not sku:
                warnings.append(f"STOCK row {excel_row}: ไม่มี ITEM ID")
                continue

            if sku in seen_skus:
                warnings.append(
                    f"STOCK row {excel_row}: ITEM ID ซ้ำใน Excel: {sku}"
                )
            seen_skus.add(sku)

            # ---------------- Unit ----------------
            unit_code = text(row.get("UNIT")) or "EA"

            try:
                with transaction.atomic():
                    unit, created = Unit.objects.get_or_create(
                        code=unit_code,
                        defaults={"name": unit_code},
                    )
                stats["units_created"] += int(created)
            except Exception as exc:
                errors.append(
                    f"STOCK row {excel_row}: Unit {unit_code}: {exc}"
                )
                continue

            # ---------------- Maker ----------------
            maker = None
            maker_name = text(row.get("MAKER"))

            if maker_name:
                try:
                    with transaction.atomic():
                        maker, created = Maker.objects.get_or_create(
                            name=maker_name
                        )
                    stats["makers_created"] += int(created)
                except Exception as exc:
                    errors.append(
                        f"STOCK row {excel_row}: Maker {maker_name}: {exc}"
                    )

            # ---------------- Supplier ----------------
            supplier = None
            vendor_name = text(row.get("Vendor"))

            if vendor_name:
                try:
                    with transaction.atomic():
                        supplier, created = get_supplier(vendor_name)
                    stats["suppliers_created"] += int(created)
                except Exception as exc:
                    errors.append(
                        f"STOCK row {excel_row}: Supplier {vendor_name}: {exc}"
                    )

            # ---------------- Location ----------------
            location = None
            location_code = text(row.get("ADDRESS"))

            if location_code:
                try:
                    with transaction.atomic():
                        location, created = Location.objects.get_or_create(
                            code=location_code,
                            defaults={
                                "name": location_code,
                                "legacy_source": "STOCK",
                                "legacy_id": location_code,
                            },
                        )
                    stats["locations_created"] += int(created)
                except Exception as exc:
                    errors.append(
                        f"STOCK row {excel_row}: Location {location_code}: {exc}"
                    )

            # ---------------- Part ----------------
            try:
                with transaction.atomic():
                    part, created = Part.objects.update_or_create(
                        sku=sku,
                        defaults={
                            "name": text(row.get("PART NAME")) or sku,
                            "description": text(row.get("PART DETAIL")),
                            "maker": maker,
                            "unit": unit,
                            "default_supplier": supplier,
                            "location": location,
                            "min_stock": decimal_value(row.get("MIN STOCK")),
                            "reorder_qty": decimal_value(row.get("TO Order")),
                            "vendor_lead_time_days": positive_int(
                                row.get("Lead Time V/D")
                            ),
                            "purchasing_lead_time_days": positive_int(
                                row.get("Lead Time Purchasing")
                            ),
                            "total_lead_time_days": positive_int(
                                row.get("Lead Time Total")
                            ),
                            "last_purchase_price": decimal_value(
                                row.get("Price")
                            ),
                            "remark": text(row.get("Remark")),
                            "image_path": text(row.get("PIC")),
                            "legacy_source": "STOCK",
                            "legacy_id": sku,
                        },
                    )
                stats["parts_processed"] += 1
                stats["parts_created"] += int(created)
            except Exception as exc:
                errors.append(
                    f"STOCK row {excel_row}: Part {sku}: {exc}"
                )
                continue

            # ---------------- Inventory ----------------
            try:
                quantity = decimal_value(row.get("REMAINING STOCK"))

                with transaction.atomic():
                    inventory, created = Inventory.objects.get_or_create(
                        part=part,
                        location=location,
                        defaults={
                            "quantity": quantity,
                            "legacy_source": "STOCK",
                            "legacy_id": sku,
                        },
                    )

                    if not created:
                        inventory.quantity = quantity
                        inventory.legacy_source = "STOCK"
                        inventory.legacy_id = sku
                        inventory.save()

                stats["inventory_processed"] += 1
                stats["inventory_created"] += int(created)
            except Exception as exc:
                errors.append(
                    f"STOCK row {excel_row}: Inventory {sku}: {exc}"
                )

            # ---------------- Part <-> Machine ----------------
            raw_machine = text(row.get("Machine Name"))
            machine_codes = split_machine_codes(raw_machine)

            if raw_machine:
                stats["machine_cells_nonempty"] += 1

            if len(machine_codes) > 1:
                stats["multi_machine_cells"] += 1

            for original_code in machine_codes:
                stats["machine_reference_occurrences"] += 1
                code_key = normalize_key(original_code)
                machine = None
                resolution_type = None

                # 1. Current code
                machine = current_by_key.get(code_key)
                if machine:
                    resolution_type = "CURRENT"

                # 2. Existing MachineCode mapping in DB
                if machine is None:
                    existing_code = MachineCode.objects.filter(
                        code__iexact=original_code
                    ).select_related("machine").first()

                    if existing_code and existing_code.machine_id:
                        machine = existing_code.machine
                        resolution_type = existing_code.code_type

                # 3. Confirmed/derived alias
                if machine is None:
                    machine = alias_target_by_key.get(code_key)

                    if machine:
                        mapping_meta = mapping_meta_by_key[code_key]
                        resolution_type = mapping_meta["code_type"]

                        try:
                            with transaction.atomic():
                                existing_alias = MachineCode.objects.filter(
                                    code__iexact=original_code
                                ).first()

                                if existing_alias:
                                    existing_alias.machine = machine
                                    existing_alias.code_type = mapping_meta[
                                        "code_type"
                                    ]
                                    existing_alias.note = mapping_meta["note"]
                                    existing_alias.active = mapping_meta[
                                        "active"
                                    ]
                                    existing_alias.save()
                                else:
                                    MachineCode.objects.create(
                                        code=original_code,
                                        machine=machine,
                                        code_type=mapping_meta["code_type"],
                                        note=mapping_meta["note"],
                                        active=mapping_meta["active"],
                                    )
                            stats["old_alias_codes_created_or_updated"] += 1
                        except Exception as exc:
                            errors.append(
                                f"STOCK row {excel_row}: Machine alias "
                                f"{original_code} -> {machine.code}: {exc}"
                            )

                # 4. Link resolved machine to part
                if machine is not None:
                    try:
                        with transaction.atomic():
                            _, created = PartMachine.objects.get_or_create(
                                part=part,
                                machine=machine,
                                defaults={
                                    "legacy_source": "STOCK",
                                    "legacy_id": f"{sku}:{original_code}",
                                },
                            )
                        stats["machine_references_resolved"] += 1
                        stats["part_machine_links_created"] += int(created)
                    except Exception as exc:
                        errors.append(
                            f"STOCK row {excel_row}: PartMachine "
                            f"{sku} -> {machine.code}: {exc}"
                        )
                    continue

                # 5. Unresolved reference: keep it, but do NOT invent a Machine.
                unresolved_occurrences[original_code] += 1
                unresolved_parts[original_code].add(sku)

                try:
                    with transaction.atomic():
                        existing_ref = MachineCode.objects.filter(
                            code__iexact=original_code
                        ).first()

                        if existing_ref is None:
                            MachineCode.objects.create(
                                code=original_code,
                                machine=None,
                                code_type="REFERENCE",
                                note=(
                                    "Found in STOCK Machine Name; "
                                    "not yet mapped to Machine Master."
                                ),
                                active=True,
                            )
                            stats["reference_codes_created"] += 1
                        elif existing_ref.machine_id is None:
                            existing_ref.code_type = "REFERENCE"
                            existing_ref.note = (
                                "Found in STOCK Machine Name; "
                                "not yet mapped to Machine Master."
                            )
                            existing_ref.active = True
                            existing_ref.save()
                except Exception as exc:
                    errors.append(
                        f"STOCK row {excel_row}: unresolved MachineCode "
                        f"{original_code}: {exc}"
                    )

        # ========================================================
        # OPTIONAL UNMATCHED CSV
        # ========================================================
        if unmatched_csv:
            try:
                with open(unmatched_csv, "w", newline="", encoding="utf-8-sig") as f:
                    writer = csv.writer(f)
                    writer.writerow(
                        [
                            "stock_code",
                            "occurrences",
                            "part_count",
                            "suggested_current_code",
                            "suggestion_reason",
                            "decision",
                            "approved_current_code",
                            "example_parts",
                            "note",
                        ]
                    )
                    for code, count in unresolved_occurrences.most_common():
                        parts = sorted(unresolved_parts[code])
                        suggestion, reason = suggested_current_code(
                            code, current_by_key
                        )
                        writer.writerow(
                            [
                                code,
                                count,
                                len(parts),
                                suggestion,
                                reason,
                                "",
                                "",
                                ", ".join(parts[:20]),
                                "",
                            ]
                        )
                self.stdout.write(
                    self.style.SUCCESS(
                        f"Unmatched machine CSV: {unmatched_csv}"
                    )
                )
            except Exception as exc:
                warnings.append(
                    f"ไม่สามารถเขียน unmatched CSV '{unmatched_csv}': {exc}"
                )

        # ========================================================
        # REPORT
        # ========================================================
        self.stdout.write("")
        self.stdout.write(self.style.MIGRATE_HEADING("=" * 60))
        self.stdout.write(
            self.style.MIGRATE_HEADING(" IMPORT PREVIEW V3.3")
        )
        self.stdout.write(self.style.MIGRATE_HEADING("=" * 60))

        report_lines = [
            ("Employees processed", stats["employees_processed"]),
            ("Employees created", stats["employees_created"]),
            ("MC NAME rows processed", stats["machine_rows_processed"]),
            ("Current machines unique", stats["current_machine_unique"]),
            ("Duplicate MC NAME rows", stats["machine_duplicate_rows"]),
            ("Machine mappings loaded", stats["machine_mappings_loaded"]),
            ("Confirmed aliases available", stats["confirmed_aliases_available"]),
            ("Parts processed", stats["parts_processed"]),
            ("Parts created", stats["parts_created"]),
            ("Inventory processed", stats["inventory_processed"]),
            ("Inventory created", stats["inventory_created"]),
            ("Machine cells non-empty", stats["machine_cells_nonempty"]),
            ("Multi-machine cells", stats["multi_machine_cells"]),
            ("Machine refs total", stats["machine_reference_occurrences"]),
            ("Machine refs resolved", stats["machine_references_resolved"]),
            (
                "Machine refs unresolved",
                sum(unresolved_occurrences.values()),
            ),
            (
                "Unresolved unique codes",
                len(unresolved_occurrences),
            ),
            ("Part ↔ Machine links created", stats["part_machine_links_created"]),
            ("Old aliases written", stats["old_alias_codes_created_or_updated"]),
            ("Reference codes created", stats["reference_codes_created"]),
            ("New Units", stats["units_created"]),
            ("New Makers", stats["makers_created"]),
            ("New Suppliers", stats["suppliers_created"]),
            ("New Locations", stats["locations_created"]),
            ("Warnings", len(warnings)),
            ("Errors", len(errors)),
        ]

        for label, value in report_lines:
            self.stdout.write(f"{label:<30}: {value}")

        if duplicate_machine_codes:
            self.stdout.write("")
            self.stdout.write(self.style.WARNING("DUPLICATE MC NAME CODES:"))
            for key in duplicate_machine_codes:
                self.stdout.write(
                    f"  - {key}: {current_code_duplicates[key]} rows"
                )

        if unresolved_occurrences:
            self.stdout.write("")
            self.stdout.write(
                self.style.WARNING(
                    f"UNMATCHED MACHINE CODES "
                    f"(showing up to {unmatched_limit}):"
                )
            )

            for code, count in unresolved_occurrences.most_common(
                unmatched_limit
            ):
                examples = ", ".join(
                    sorted(unresolved_parts[code])[:5]
                )
                suggestion, _ = suggested_current_code(
                    code, current_by_key
                )
                suggestion_text = (
                    f" | suggest: {suggestion} (NOT auto-mapped)"
                    if suggestion else ""
                )
                self.stdout.write(
                    f"  - {code}: {count} refs "
                    f"| parts: {examples}{suggestion_text}"
                )

            if len(unresolved_occurrences) > unmatched_limit:
                self.stdout.write(
                    f"  ... และอีก "
                    f"{len(unresolved_occurrences) - unmatched_limit} codes"
                )

        if warnings:
            self.stdout.write("")
            self.stdout.write(self.style.WARNING("WARNINGS:"))
            for warning in warnings[:50]:
                self.stdout.write(f"  - {warning}")
            if len(warnings) > 50:
                self.stdout.write(
                    f"  ... และอีก {len(warnings) - 50} warnings"
                )

        if errors:
            self.stdout.write("")
            self.stdout.write(self.style.ERROR("ERRORS:"))
            for error in errors[:100]:
                self.stdout.write(f"  - {error}")

            transaction.set_rollback(True)
            raise CommandError(
                f"Import aborted because {len(errors)} error(s) were found."
            )

        # ========================================================
        # DRY RUN / COMMIT
        # ========================================================
        if dry_run:
            transaction.set_rollback(True)
            self.stdout.write("")
            self.stdout.write(self.style.WARNING("=" * 60))
            self.stdout.write(
                self.style.WARNING("DRY RUN COMPLETE - DATABASE ROLLED BACK")
            )
            self.stdout.write(self.style.WARNING("=" * 60))
            return

        self.stdout.write("")
        self.stdout.write(self.style.SUCCESS("=" * 60))
        self.stdout.write(
            self.style.SUCCESS("IMPORT COMPLETE - SAVED TO DATABASE")
        )
        self.stdout.write(self.style.SUCCESS("=" * 60))
