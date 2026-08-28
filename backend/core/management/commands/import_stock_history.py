import csv
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook

from core.models import Employee, Machine, MachineCode, Part, StockTransaction


def txt(value):
    return "" if value is None else str(value).strip()


def dec(value):
    if value in (None, ""):
        return Decimal("0")
    try:
        return Decimal(str(value))
    except (InvalidOperation, TypeError, ValueError):
        return Decimal("0")


def fields(model):
    return {f.name: f for f in model._meta.get_fields()}


def first_field(model, candidates):
    names = fields(model)
    for name in candidates:
        if name in names:
            return name
    return None


def sheet_rows(ws):
    headers = [txt(c.value) for c in ws[1]]
    result = []

    for row_no, values in enumerate(
        ws.iter_rows(min_row=2, values_only=True), start=2
    ):
        if not any(v not in (None, "") for v in values):
            continue

        row = {}
        for i, header in enumerate(headers):
            if header:
                row[header] = values[i] if i < len(values) else None

        row["_row"] = row_no
        result.append(row)

    return result


def parse_date(value):
    if value in (None, ""):
        return None

    if isinstance(value, datetime):
        y, m, d = value.year, value.month, value.day
    elif isinstance(value, date):
        y, m, d = value.year, value.month, value.day
    else:
        s = txt(value)
        parsed = None
        for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
            try:
                parsed = datetime.strptime(s, fmt)
                break
            except ValueError:
                pass
        if parsed is None:
            return None
        y, m, d = parsed.year, parsed.month, parsed.day

    if y >= 2400:
        y -= 543

    try:
        return date(y, m, d)
    except ValueError:
        return None


def parse_time(value):
    if value in (None, ""):
        return time(0, 0)

    if isinstance(value, datetime):
        return value.time().replace(microsecond=0)

    if isinstance(value, time):
        return value.replace(microsecond=0)

    s = txt(value)
    for fmt in ("%H:%M:%S", "%H:%M"):
        try:
            return datetime.strptime(s, fmt).time()
        except ValueError:
            pass

    return time(0, 0)


def parse_datetime(date_value, time_value):
    d = parse_date(date_value)
    if d is None:
        return None

    dt = datetime.combine(d, parse_time(time_value))
    if timezone.is_naive(dt):
        dt = timezone.make_aware(dt, timezone.get_current_timezone())
    return dt


def find_part(item_id):
    item_id = txt(item_id)
    if not item_id:
        return None

    field = first_field(Part, ("sku", "item_id", "code"))
    if not field:
        raise CommandError("Part model has no sku/item_id/code field")

    return Part.objects.filter(**{f"{field}__iexact": item_id}).first()


def find_employee(value):
    value = txt(value)
    if not value:
        return None

    if "employee_code" in fields(Employee):
        obj = Employee.objects.filter(employee_code__iexact=value).first()
        if obj:
            return obj

    if "name" in fields(Employee):
        obj = Employee.objects.filter(name__iexact=value).first()
        if obj:
            return obj

        candidates = list(
            Employee.objects.filter(name__istartswith=value).order_by("name")[:2]
        )
        if len(candidates) == 1:
            return candidates[0]

    return None


def find_machine(value):
    value = txt(value)
    if not value:
        return None

    if "code" in fields(Machine):
        obj = Machine.objects.filter(code__iexact=value).first()
        if obj:
            return obj

    code = (
        MachineCode.objects
        .filter(code__iexact=value, active=True)
        .select_related("machine")
        .first()
    )
    if code and code.machine_id:
        return code.machine

    return None


def choice_value(model, field_name, preferred, fallback):
    field = fields(model).get(field_name)
    if not field:
        return fallback

    choices = dict(getattr(field, "choices", []) or [])
    if not choices:
        return fallback

    for value in preferred:
        if value in choices:
            return value

    lower_words = [str(v).lower() for v in preferred]
    for value, label in choices.items():
        label_lower = str(label).lower()
        if any(word in label_lower for word in lower_words):
            return value

    return fallback


def transaction_type(kind):
    if kind == "RECEIVE":
        preferred = ("IN", "RECEIVE", "RECEIPT", "receive", "in")
        fallback = "IN"
    else:
        preferred = ("OUT", "ISSUE", "WITHDRAW", "withdraw", "issue", "out")
        fallback = "OUT"

    return choice_value(
        StockTransaction,
        "transaction_type",
        preferred,
        fallback,
    )


def reference_type(kind):
    preferred = (
        "LEGACY",
        "EXCEL",
        kind,
        kind.lower(),
        "MANUAL",
        "OTHER",
    )

    field = fields(StockTransaction).get("reference_type")
    choices = dict(getattr(field, "choices", []) or []) if field else {}

    if choices:
        for candidate in preferred:
            if candidate in choices:
                return candidate
        return next(iter(choices.keys()))

    return kind


def transaction_no(kind, legacy_id):
    prefix = "RCV" if kind == "RECEIVE" else "WDR"
    return f"{prefix}-{legacy_id}"


def legacy_exists(kind, legacy_id):
    model_fields = fields(StockTransaction)

    if "legacy_source" in model_fields and "legacy_id" in model_fields:
        return StockTransaction.objects.filter(
            legacy_source=kind,
            legacy_id=legacy_id,
        ).exists()

    if "transaction_no" in model_fields:
        return StockTransaction.objects.filter(
            transaction_no=transaction_no(kind, legacy_id)
        ).exists()

    return False


def build_data(kind, row, part, machine, employee):
    key_col = "RECEIVE KEY" if kind == "RECEIVE" else "WITHDRAW KEY"
    legacy_id = txt(row.get(key_col))
    occurred_at = parse_datetime(row.get("DATE"), row.get("TIME"))

    if not occurred_at:
        raise ValueError("invalid DATE/TIME")

    amount = dec(row.get("AMOUNT"))
    if amount <= 0:
        raise ValueError("AMOUNT must be > 0")

    data = {
        "legacy_source": kind,
        "legacy_id": legacy_id,
        "transaction_no": transaction_no(kind, legacy_id),
        "part": part,
        "location": getattr(part, "location", None),
        "transaction_type": transaction_type(kind),
        "quantity": amount,
        "machine": machine,
        "employee": employee,
        "reference_type": reference_type(kind),
        "reference_id": legacy_id,
        "transaction_date": occurred_at,
        "remark": "",
        "created_by": None,
    }

    remark_parts = []

    original = txt(row.get("REMARK"))
    if original:
        remark_parts.append(original)

    raw_recorder = txt(row.get("RECORDED BY"))
    if raw_recorder:
        remark_parts.append(f"Legacy RECORDED BY: {raw_recorder}")

    if kind == "WITHDRAW":
        raw_withdraw_by = txt(row.get("WITHDRAW BY"))
        if raw_withdraw_by and employee is None:
            remark_parts.append(f"Legacy WITHDRAW BY: {raw_withdraw_by}")

    raw_machine = txt(row.get("MACHINE NAME"))
    if raw_machine and machine is None:
        remark_parts.append(f"Legacy MACHINE NAME: {raw_machine}")

    data["remark"] = " | ".join(remark_parts)
    return data


class Command(BaseCommand):
    help = (
        "Import historical RECEIVE/WITHDRAW rows into StockTransaction "
        "without changing current Inventory."
    )

    def add_arguments(self, parser):
        parser.add_argument("xlsx", help="Path to ROD MM STOCK.xlsx")
        parser.add_argument(
            "--dry-run",
            action="store_true",
            help="Validate/import then rollback all database changes.",
        )
        parser.add_argument(
            "--report",
            default="history_import_report_final.csv",
            help="CSV report path.",
        )

    @transaction.atomic
    def handle(self, *args, **options):
        xlsx = Path(options["xlsx"])
        dry_run = options["dry_run"]
        report_path = Path(options["report"])

        if not xlsx.exists():
            raise CommandError(f"Excel not found: {xlsx}")

        wb = load_workbook(xlsx, data_only=True, read_only=True)

        for required in ("RECEIVE", "WITHDRAW"):
            if required not in wb.sheetnames:
                raise CommandError(f"Missing sheet: {required}")

        required_tx_fields = {
            "legacy_source",
            "legacy_id",
            "transaction_no",
            "part",
            "location",
            "transaction_type",
            "quantity",
            "machine",
            "employee",
            "reference_type",
            "reference_id",
            "transaction_date",
            "remark",
            "created_by",
        }

        missing = sorted(required_tx_fields - set(fields(StockTransaction)))
        if missing:
            raise CommandError(
                "StockTransaction missing fields: " + ", ".join(missing)
            )

        stats = {
            "receive_rows": 0,
            "withdraw_rows": 0,
            "imported": 0,
            "duplicates": 0,
            "skipped_zero_amount": 0,
            "skipped_missing_part": 0,
            "missing_machine_warning": 0,
            "missing_employee_warning": 0,
            "errors": 0,
        }

        report = []

        for kind in ("RECEIVE", "WITHDRAW"):
            rows = sheet_rows(wb[kind])
            stats["receive_rows" if kind == "RECEIVE" else "withdraw_rows"] = len(rows)
            self.stdout.write(f"{kind}: {len(rows)} rows")

            for row in rows:
                excel_row = row["_row"]
                item_id = txt(row.get("ITEM ID"))
                key_col = "RECEIVE KEY" if kind == "RECEIVE" else "WITHDRAW KEY"
                legacy_id = txt(row.get(key_col))

                if not legacy_id:
                    stats["errors"] += 1
                    report.append(
                        [kind, excel_row, item_id, "", "ERROR", "missing legacy key"]
                    )
                    continue

                if legacy_exists(kind, legacy_id):
                    stats["duplicates"] += 1
                    report.append(
                        [kind, excel_row, item_id, legacy_id, "DUPLICATE", "already imported"]
                    )
                    continue

                amount = dec(row.get("AMOUNT"))
                if amount <= 0:
                    stats["skipped_zero_amount"] += 1
                    report.append(
                        [
                            kind,
                            excel_row,
                            item_id,
                            legacy_id,
                            "SKIPPED_ZERO_AMOUNT",
                            f"AMOUNT={txt(row.get('AMOUNT')) or '0'}",
                        ]
                    )
                    continue

                part = find_part(item_id)
                if not part:
                    stats["skipped_missing_part"] += 1
                    report.append(
                        [
                            kind,
                            excel_row,
                            item_id,
                            legacy_id,
                            "SKIPPED_MISSING_PART",
                            "part not found in current Part master",
                        ]
                    )
                    continue

                raw_machine = txt(row.get("MACHINE NAME"))
                machine = find_machine(raw_machine)
                if raw_machine and machine is None:
                    stats["missing_machine_warning"] += 1

                employee = None
                if kind == "WITHDRAW":
                    raw_employee = txt(row.get("WITHDRAW BY"))
                    employee = find_employee(raw_employee)
                    if raw_employee and employee is None:
                        stats["missing_employee_warning"] += 1

                try:
                    data = build_data(kind, row, part, machine, employee)

                    with transaction.atomic():
                        StockTransaction.objects.create(**data)

                    stats["imported"] += 1

                    warnings = []
                    if raw_machine and machine is None:
                        warnings.append("machine unmatched")
                    if (
                        kind == "WITHDRAW"
                        and txt(row.get("WITHDRAW BY"))
                        and employee is None
                    ):
                        warnings.append("employee unmatched")

                    if warnings:
                        report.append(
                            [
                                kind,
                                excel_row,
                                item_id,
                                legacy_id,
                                "IMPORTED_WITH_WARNING",
                                "; ".join(warnings),
                            ]
                        )

                except Exception as exc:
                    stats["errors"] += 1
                    report.append(
                        [kind, excel_row, item_id, legacy_id, "ERROR", str(exc)]
                    )

        report_path.parent.mkdir(parents=True, exist_ok=True)
        with report_path.open("w", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f)
            writer.writerow(
                ["sheet", "excel_row", "item_id", "legacy_id", "result", "detail"]
            )
            writer.writerows(report)

        total_rows = stats["receive_rows"] + stats["withdraw_rows"]
        accounted = (
            stats["imported"]
            + stats["duplicates"]
            + stats["skipped_zero_amount"]
            + stats["skipped_missing_part"]
            + stats["errors"]
        )

        self.stdout.write("")
        self.stdout.write(self.style.MIGRATE_HEADING("Import summary"))
        for key, value in stats.items():
            self.stdout.write(f"{key}: {value}")

        self.stdout.write(f"total_rows: {total_rows}")
        self.stdout.write(f"accounted_rows: {accounted}")
        self.stdout.write(f"Report: {report_path}")

        if total_rows != accounted:
            raise CommandError(
                f"Safety check failed: total_rows={total_rows}, accounted_rows={accounted}"
            )

        if dry_run:
            transaction.set_rollback(True)
            self.stdout.write(
                self.style.WARNING(
                    "DRY RUN: all StockTransaction changes rolled back. "
                    "Inventory was never modified."
                )
            )
        else:
            if stats["errors"] > 0:
                transaction.set_rollback(True)
                raise CommandError(
                    f"Real import aborted because errors={stats['errors']}. "
                    "No StockTransaction rows were committed."
                )

            self.stdout.write(
                self.style.SUCCESS(
                    "History import completed successfully. "
                    "Inventory was NOT modified."
                )
            )
